from docx import Document
from docx.shared import Pt, RGBColor, Inches, Cm
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_ALIGN_VERTICAL, WD_TABLE_ALIGNMENT
from docx.oxml.ns import qn
from docx.oxml import OxmlElement
import copy

doc = Document()

# ── Page margins ──────────────────────────────────────────────
for section in doc.sections:
    section.top_margin    = Inches(1)
    section.bottom_margin = Inches(1)
    section.left_margin   = Inches(1)
    section.right_margin  = Inches(1)

# ── Helpers ───────────────────────────────────────────────────
def set_font(run, bold=False, size=11, color=None, italic=False):
    run.bold   = bold
    run.italic = italic
    run.font.size = Pt(size)
    if color:
        run.font.color.rgb = RGBColor(*color)

def add_paragraph(doc, text="", bold=False, size=11, color=None,
                  align=WD_ALIGN_PARAGRAPH.LEFT, space_before=0, space_after=4):
    p = doc.add_paragraph()
    p.alignment = align
    p.paragraph_format.space_before = Pt(space_before)
    p.paragraph_format.space_after  = Pt(space_after)
    if text:
        run = p.add_run(text)
        set_font(run, bold=bold, size=size, color=color)
    return p

def add_bullet(doc, text, bold_prefix=None, size=10.5):
    p = doc.add_paragraph(style="List Bullet")
    p.paragraph_format.space_after  = Pt(2)
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.left_indent  = Inches(0.25)
    if bold_prefix:
        r1 = p.add_run(bold_prefix)
        set_font(r1, bold=True, size=size)
        r2 = p.add_run(text)
        set_font(r2, size=size)
    else:
        r = p.add_run(text)
        set_font(r, size=size)
    return p

def shade_cell(cell, hex_color):
    tc   = cell._tc
    tcPr = tc.get_or_add_tcPr()
    shd  = OxmlElement("w:shd")
    shd.set(qn("w:val"),   "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"),  hex_color)
    tcPr.append(shd)

def set_cell_border(cell, **kwargs):
    tc   = cell._tc
    tcPr = tc.get_or_add_tcPr()
    tcBorders = OxmlElement("w:tcBorders")
    for edge in ("top","left","bottom","right","insideH","insideV"):
        if edge in kwargs:
            tag = OxmlElement(f"w:{edge}")
            for k, v in kwargs[edge].items():
                tag.set(qn(f"w:{k}"), v)
            tcBorders.append(tag)
    tcPr.append(tcBorders)

def section_header(doc, title, dark=True):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(10)
    p.paragraph_format.space_after  = Pt(4)
    run = p.add_run(title.upper())
    run.bold = True
    run.font.size = Pt(10)
    if dark:
        run.font.color.rgb = RGBColor(255, 255, 255)
        # paragraph shading
        pPr  = p._p.get_or_add_pPr()
        shd  = OxmlElement("w:shd")
        shd.set(qn("w:val"),   "clear")
        shd.set(qn("w:color"), "auto")
        shd.set(qn("w:fill"),  "202124")
        pPr.append(shd)
    else:
        run.font.color.rgb = RGBColor(32, 33, 36)
        pPr  = p._p.get_or_add_pPr()
        shd  = OxmlElement("w:shd")
        shd.set(qn("w:val"),   "clear")
        shd.set(qn("w:color"), "auto")
        shd.set(qn("w:fill"),  "F5F5F5")
        pPr.append(shd)
        # left border accent
        pBdr = OxmlElement("w:pBdr")
        left = OxmlElement("w:left")
        left.set(qn("w:val"),   "single")
        left.set(qn("w:sz"),    "24")
        left.set(qn("w:space"), "4")
        left.set(qn("w:color"), "202124")
        pBdr.append(left)
        pPr.append(pBdr)
    return p

def add_sub_label(doc, text):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(6)
    p.paragraph_format.space_after  = Pt(2)
    r = p.add_run(text)
    r.bold = True
    r.font.size = Pt(10.5)
    return p

def add_hr(doc):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(4)
    p.paragraph_format.space_after  = Pt(4)
    pPr  = p._p.get_or_add_pPr()
    pBdr = OxmlElement("w:pBdr")
    bot  = OxmlElement("w:bottom")
    bot.set(qn("w:val"),   "single")
    bot.set(qn("w:sz"),    "4")
    bot.set(qn("w:space"), "1")
    bot.set(qn("w:color"), "CCCCCC")
    pBdr.append(bot)
    pPr.append(pBdr)

# ═══════════════════════════════════════════════════════════════
# HEADER
# ═══════════════════════════════════════════════════════════════
p = add_paragraph(doc, "FURNITURE & THEN SOME",
                  bold=True, size=16,
                  align=WD_ALIGN_PARAGRAPH.CENTER, space_after=2)

p = add_paragraph(doc, "Wallpaper Installer — Expectations & Pay Sheet",
                  size=11, color=(80, 80, 80),
                  align=WD_ALIGN_PARAGRAPH.CENTER, space_after=4)

p = doc.add_paragraph()
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
p.paragraph_format.space_after = Pt(2)
r = p.add_run("Position: ")
set_font(r, bold=True, size=10.5)
r = p.add_run("Wallpaper Installer     ")
set_font(r, size=10.5)
r = p.add_run("Compensation: ")
set_font(r, bold=True, size=10.5)
r = p.add_run("Per Job (Piece Rate)")
set_font(r, size=10.5)

add_hr(doc)

# ═══════════════════════════════════════════════════════════════
# PAY STRUCTURE TABLE
# ═══════════════════════════════════════════════════════════════
section_header(doc, "Pay Structure — Installer Rate")

tbl = doc.add_table(rows=3, cols=2)
tbl.style = "Table Grid"
tbl.alignment = WD_TABLE_ALIGNMENT.LEFT

col_widths = [Inches(3.5), Inches(2.5)]
for row in tbl.rows:
    for i, cell in enumerate(row.cells):
        cell.width = col_widths[i]

# Header row
hdr_cells = tbl.rows[0].cells
for cell in hdr_cells:
    shade_cell(cell, "F5F5F5")

def cell_text(cell, text, bold=False, size=10.5, color=None):
    cell.text = ""
    p = cell.paragraphs[0]
    p.paragraph_format.space_before = Pt(3)
    p.paragraph_format.space_after  = Pt(3)
    r = p.add_run(text)
    set_font(r, bold=bold, size=size, color=color)

cell_text(hdr_cells[0], "Service Type", bold=True)
cell_text(hdr_cells[1], "Rate per Sq Ft", bold=True)

data = [
    ("Peel & Stick Wallpaper", "$1.75 – $2.25"),
    ("Pasted Wallpaper",       "$2.00 – $2.75"),
]
for i, (svc, rate) in enumerate(data):
    row = tbl.rows[i + 1]
    cell_text(row.cells[0], svc)
    cell_text(row.cells[1], rate, bold=True)

p = doc.add_paragraph()
r = p.add_run("Rate is determined by experience, speed, and quality of work.")
set_font(r, size=9, color=(100, 100, 100), italic=True)
p.paragraph_format.space_before = Pt(4)
p.paragraph_format.space_after  = Pt(4)

# ═══════════════════════════════════════════════════════════════
# PAY GROWTH SYSTEM TABLE
# ═══════════════════════════════════════════════════════════════
section_header(doc, "Pay Growth System")

tbl2 = doc.add_table(rows=4, cols=3)
tbl2.style = "Table Grid"

col_widths2 = [Inches(1.5), Inches(1.5), Inches(3.0)]
for row in tbl2.rows:
    for i, cell in enumerate(row.cells):
        cell.width = col_widths2[i]

hdr2 = tbl2.rows[0].cells
for cell in hdr2:
    shade_cell(cell, "F5F5F5")

cell_text(hdr2[0], "Level",         bold=True)
cell_text(hdr2[1], "Rate / Sq Ft",  bold=True)
cell_text(hdr2[2], "Criteria",      bold=True)

growth_data = [
    ("Training Phase",     "$1.25/sq ft",         "Starting rate — all new installers begin here"),
    ("Standard Installer", "$1.75 – $2.25",      "Consistent quality and reliability"),
    ("Top Performer",      "$2.50 – $2.75\n+ bonuses", "Speed, clean installs, reviews, leadership"),
]
for i, (lvl, rate, crit) in enumerate(growth_data):
    row = tbl2.rows[i + 1]
    cell_text(row.cells[0], lvl)
    cell_text(row.cells[1], rate, bold=True)
    cell_text(row.cells[2], crit)

# Plain-language starting pay notice
p_notice = doc.add_paragraph()
p_notice.paragraph_format.space_before = Pt(10)
p_notice.paragraph_format.space_after  = Pt(6)
r_notice = p_notice.add_run(
    "All new installers start at the Training Pay rate of $1.25 per square foot. "
    "This is your starting pay while you learn our standards and processes. "
    "As your skills, speed, and quality improve, your pay rate will increase accordingly."
)
set_font(r_notice, size=10.5)

# Bonus callout (shaded green)
p_bonus = doc.add_paragraph()
p_bonus.paragraph_format.space_before = Pt(6)
p_bonus.paragraph_format.space_after  = Pt(2)
r = p_bonus.add_run("Raises & Bonuses")
set_font(r, bold=True, size=10.5)

# Green left border paragraph style
def green_left_border(p):
    pPr  = p._p.get_or_add_pPr()
    shd  = OxmlElement("w:shd")
    shd.set(qn("w:val"),   "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"),  "F0F7F0")
    pPr.append(shd)
    pBdr = OxmlElement("w:pBdr")
    left = OxmlElement("w:left")
    left.set(qn("w:val"),   "single")
    left.set(qn("w:sz"),    "24")
    left.set(qn("w:space"), "8")
    left.set(qn("w:color"), "2E7D32")
    pBdr.append(left)
    pPr.append(pBdr)

green_left_border(p_bonus)

bonus_items = [
    ("Raises are earned, not given by time. ", "Rate increases by +$0.25/sq ft when you consistently hit speed targets and maintain clean installs with no callbacks."),
    ("$25 – $50 per 5-star review ", "when the client mentions you by name."),
    ("Bonuses ", "for fast & clean installs, no callbacks (30+ days), and leading jobs independently."),
]
for bold_part, rest in bonus_items:
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after  = Pt(2)
    p.paragraph_format.left_indent  = Inches(0.25)
    green_left_border(p)
    r1 = p.add_run("• " + bold_part)
    set_font(r1, bold=True, size=10.5)
    r2 = p.add_run(rest)
    set_font(r2, size=10.5)

# ═══════════════════════════════════════════════════════════════
# EXPECTATIONS
# ═══════════════════════════════════════════════════════════════
section_header(doc, "Expectations (Non-Negotiable)")

expectations = {
    "Professionalism": [
        ("On time means ", "10 minutes early."),
        ("Respect ", "every client and their home."),
        (None, "Maintain a clean, professional appearance."),
    ],
    "Work Quality": [
        (None, "Straight seams, no bubbles, no lifting."),
        ("Precision over speed ", "— consistency first."),
    ],
    "Clean Work Area": [
        (None, "No mess left behind."),
        (None, "Protect floors, furniture, and walls."),
        (None, "Full cleanup after every job."),
    ],
    "Efficiency": [
        (None, "Stay on task — no excessive phone use."),
        (None, "Hit expected completion timelines."),
    ],
}

for label, items in expectations.items():
    add_sub_label(doc, label)
    for bold_part, rest in items:
        p = doc.add_paragraph()
        p.paragraph_format.space_after  = Pt(2)
        p.paragraph_format.space_before = Pt(0)
        p.paragraph_format.left_indent  = Inches(0.25)
        if bold_part:
            r1 = p.add_run("• " + bold_part)
            set_font(r1, bold=True, size=10.5)
            r2 = p.add_run(rest)
            set_font(r2, size=10.5)
        else:
            r = p.add_run("• " + rest)
            set_font(r, size=10.5)

# ═══════════════════════════════════════════════════════════════
# IMMEDIATE TERMINATION
# ═══════════════════════════════════════════════════════════════
section_header(doc, "Immediate Termination")

termination_items = [
    "Discussing pricing with clients",
    "No call / no show",
    "Repeated lateness",
    "Poor quality work",
    "Disrespect to clients or team members",
    "Damaging property due to carelessness",
    "Taking side jobs from company clients",
    "Showing up under the influence",
]

def red_left_border(p):
    pPr  = p._p.get_or_add_pPr()
    shd  = OxmlElement("w:shd")
    shd.set(qn("w:val"),   "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"),  "FFF8F8")
    pPr.append(shd)
    pBdr = OxmlElement("w:pBdr")
    left = OxmlElement("w:left")
    left.set(qn("w:val"),   "single")
    left.set(qn("w:sz"),    "24")
    left.set(qn("w:space"), "8")
    left.set(qn("w:color"), "CC0000")
    pBdr.append(left)
    pPr.append(pBdr)

for item in termination_items:
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after  = Pt(2)
    p.paragraph_format.left_indent  = Inches(0.25)
    red_left_border(p)
    r = p.add_run("• " + item)
    set_font(r, size=10.5)

# ═══════════════════════════════════════════════════════════════
# COMMUNICATION RULES
# ═══════════════════════════════════════════════════════════════
section_header(doc, "Communication Rules")

comm_items = [
    (None, "Confirm jobs the day before."),
    (None, "Notify the team immediately if running late."),
    (None, "Report any issues or concerns immediately."),
    ("Do not ", "quote prices or negotiate with clients directly."),
]
for bold_part, rest in comm_items:
    p = doc.add_paragraph()
    p.paragraph_format.space_after  = Pt(2)
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.left_indent  = Inches(0.25)
    if bold_part:
        r1 = p.add_run("• " + bold_part)
        set_font(r1, bold=True, size=10.5)
        r2 = p.add_run(rest)
        set_font(r2, size=10.5)
    else:
        r = p.add_run("• " + rest)
        set_font(r, size=10.5)

# ═══════════════════════════════════════════════════════════════
# WHAT SUCCESS LOOKS LIKE  |  GROWTH OPPORTUNITY  (side by side via table)
# ═══════════════════════════════════════════════════════════════
p_space = doc.add_paragraph()
p_space.paragraph_format.space_before = Pt(10)
p_space.paragraph_format.space_after  = Pt(0)

tbl3 = doc.add_table(rows=1, cols=2)
tbl3.style = "Table Grid"
tbl3.alignment = WD_TABLE_ALIGNMENT.LEFT

left_cell  = tbl3.rows[0].cells[0]
right_cell = tbl3.rows[0].cells[1]
left_cell.width  = Inches(3.0)
right_cell.width = Inches(3.0)

def fill_side_cell(cell, title, items):
    cell.text = ""
    # Title paragraph
    p = cell.paragraphs[0]
    p.paragraph_format.space_before = Pt(4)
    p.paragraph_format.space_after  = Pt(4)
    pPr = p._p.get_or_add_pPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"),   "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"),  "F5F5F5")
    pPr.append(shd)
    pBdr = OxmlElement("w:pBdr")
    left = OxmlElement("w:left")
    left.set(qn("w:val"),   "single")
    left.set(qn("w:sz"),    "24")
    left.set(qn("w:space"), "4")
    left.set(qn("w:color"), "202124")
    pBdr.append(left)
    pPr.append(pBdr)
    r = p.add_run(title.upper())
    set_font(r, bold=True, size=10)

    for item in items:
        p2 = cell.add_paragraph()
        p2.paragraph_format.space_before = Pt(0)
        p2.paragraph_format.space_after  = Pt(2)
        p2.paragraph_format.left_indent  = Inches(0.15)
        r2 = p2.add_run("• " + item)
        set_font(r2, size=10.5)

fill_side_cell(left_cell, "What Success Looks Like", [
    "You show up without being chased",
    "Your work is clean and consistent",
    "Clients trust you in their home",
    "You move with urgency and pride",
    "You represent the brand at a high level",
])

fill_side_cell(right_cell, "Growth Opportunity", [
    "Lead installer roles",
    "Higher-paying jobs",
    "Consistent weekly work",
    "Long-term growth with the company",
])

# ═══════════════════════════════════════════════════════════════
# AGREEMENT / SIGNATURE
# ═══════════════════════════════════════════════════════════════
section_header(doc, "Agreement")

p = doc.add_paragraph(
    "I understand the expectations, standards, and pay structure listed above "
    "and agree to uphold them as a representative of Furniture & Then Some."
)
p.paragraph_format.space_after = Pt(14)
for run in p.runs:
    set_font(run, size=10.5)

sig_fields = [
    ("Name:",      "_" * 40),
    ("Signature:", "_" * 40),
    ("Date:",      "_" * 25),
]
for label, line in sig_fields:
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(4)
    p.paragraph_format.space_after  = Pt(4)
    r1 = p.add_run(label + "  ")
    set_font(r1, bold=True, size=10.5)
    r2 = p.add_run(line)
    set_font(r2, size=10.5)

# ═══════════════════════════════════════════════════════════════
# FOOTER
# ═══════════════════════════════════════════════════════════════
add_hr(doc)
p = add_paragraph(doc, "Furniture & Then Some",
                  bold=True, size=11,
                  align=WD_ALIGN_PARAGRAPH.CENTER, space_after=2)
p = add_paragraph(doc, "We grow together — welcome to the team.",
                  size=10, color=(80, 80, 80),
                  align=WD_ALIGN_PARAGRAPH.CENTER)

doc.save("/home/user/successmarcus31/wallpaper-installer-expectations.docx")
print("Done.")
